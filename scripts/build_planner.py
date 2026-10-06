"""Builds the Budget & Finance Planner spreadsheet (Excel / Google Sheets compatible).

Usage: python3 scripts/build_planner.py [output_path] [--demo]
  Or import it: build(out, year=2027, currency="$", theme="sage", demo=False)
  --demo fills a full year of sample transactions (used for listing screenshots only).
"""
import sys
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

YEAR = 2027
TX_ROWS = 1000  # transaction rows available to the buyer

# Palette (PRIMARY / PRIMARY_LIGHT / CREAM / ACCENT / BAR are set per theme by _apply_theme)
THEMES = {
    "sage": {"primary": "2F6B4F", "light": "E3F0E8", "cream": "FAF7F0", "accent": "C9A227", "bar": "7FB394"},
    "blush": {"primary": "9E4A5F", "light": "F6E4E8", "cream": "FFF8F6", "accent": "D4A373", "bar": "D99AAB"},
    "midnight": {"primary": "1F3A5F", "light": "E2E9F3", "cream": "F7F8FB", "accent": "E0A458", "bar": "7D9CC4"},
}
CURRENCY_WORDS = {"$": "dollar", "£": "pound", "€": "euro"}
PRIMARY = PRIMARY_LIGHT = CREAM = ACCENT = BAR = ""
INK = "1F2A24"
MUTED = "6B7A71"
RED_LIGHT = "F8D7D3"
RED = "B3412F"
WHITE = "FFFFFF"

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
INCOME_CATS = ["Salary", "Side Hustle", "Freelance", "Investments", "Gifts", "Other Income"]
EXPENSE_CATS = [
    ("Rent / Mortgage", 1200), ("Utilities", 180), ("Groceries", 450), ("Transport", 200),
    ("Insurance", 150), ("Phone & Internet", 90), ("Dining Out", 150), ("Entertainment", 80),
    ("Subscriptions", 40), ("Shopping", 120), ("Health", 60), ("Personal Care", 50),
    ("Education", 30), ("Gifts & Donations", 50), ("Travel", 100), ("Debt Payments", 250),
    ("Savings Transfer", 300), ("Miscellaneous", 50),
]

thin = Side(style="thin", color="D9D4C7")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEADER_FONT = Font(name="Calibri", bold=True, color=WHITE, size=11)
INPUT_FILL = PatternFill("solid", fgColor=WHITE)
HEADER_FILL = BG_FILL = SOFT_FILL = None
MONEY = CURRENCY_WORD = ""
PCT = "0.0%"


def _apply_theme(theme, currency):
    global PRIMARY, PRIMARY_LIGHT, CREAM, ACCENT, BAR, HEADER_FILL, BG_FILL, SOFT_FILL, MONEY, CURRENCY_WORD
    t = THEMES[theme]
    PRIMARY, PRIMARY_LIGHT, CREAM, ACCENT, BAR = t["primary"], t["light"], t["cream"], t["accent"], t["bar"]
    HEADER_FILL = PatternFill("solid", fgColor=PRIMARY)
    BG_FILL = PatternFill("solid", fgColor=CREAM)
    SOFT_FILL = PatternFill("solid", fgColor=PRIMARY_LIGHT)
    MONEY = f'"{currency}"#,##0.00;[Red]-"{currency}"#,##0.00'
    CURRENCY_WORD = CURRENCY_WORDS.get(currency, "dollar")


def base(ws, widths, tab_color=None):
    ws.sheet_properties.tabColor = tab_color or PRIMARY
    ws.sheet_view.showGridLines = False
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=1, max_row=60, min_col=1, max_col=20):
        for c in row:
            c.fill = BG_FILL


def title(ws, text, subtitle, width_cols="B1:H1"):
    ws["B1"] = text
    ws["B1"].font = Font(name="Georgia", size=22, bold=True, color=PRIMARY)
    ws["B2"] = subtitle
    ws["B2"].font = Font(name="Calibri", size=11, italic=True, color=MUTED)
    ws.row_dimensions[1].height = 36


def header(ws, row, col, labels):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=label)
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 22


def cell(ws, ref, value=None, fmt=None, bold=False, fill=None, color=INK, align=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = Font(name="Calibri", bold=bold, color=color)
    c.border = BORDER
    c.fill = fill or INPUT_FILL
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = Alignment(horizontal=align)
    return c


def build(out, year=2027, currency="$", theme="sage", demo=False):
    """Write the planner workbook to `out`. Returns the output path."""
    global YEAR
    YEAR = year
    _apply_theme(theme, currency)

    wb = Workbook()

    # ---------------------------------------------------------------- Start Here
    ws = wb.active
    ws.title = "Start Here"
    base(ws, {"A": 3, "B": 100})
    title(ws, f"Your {YEAR} Budget & Finance Planner", f"Take control of every {CURRENCY_WORD} - in 10 minutes a week.")
    steps = [
        ("1. SETUP", "Open the 'Setup' tab. Enter your starting balance and your monthly budget for each "
                     "category. Rename any category to fit your life - the whole workbook updates automatically."),
        ("2. LOG", "Each time you earn or spend, add a row on the 'Transactions' tab: date, description, "
                   "pick a category from the dropdown, and the amount. The Type column fills itself in."),
        ("3. REVIEW", "Check the 'Dashboard' any time: income vs. spending per month, savings rate, "
                      "and which categories are over budget (highlighted red)."),
        ("4. GROW", "Track big goals on 'Savings Goals' (it tells you how much to save per month) and plan "
                    "your way out of debt on 'Debt Payoff' (it tells you how many months until you're free)."),
        ("TIPS", "Only edit WHITE cells. Shaded cells contain formulas. Works in Microsoft Excel, "
                 "Google Sheets (File > Import), Apple Numbers and LibreOffice. Make a copy each year!"),
    ]
    r = 4
    for head, body in steps:
        ws.cell(row=r, column=2, value=head).font = Font(name="Georgia", size=13, bold=True, color=PRIMARY)
        b = ws.cell(row=r + 1, column=2, value=body)
        b.font = Font(name="Calibri", size=11, color=INK)
        b.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r + 1].height = 34
        r += 3
    ws.cell(row=r + 1, column=2, value="Thank you for your purchase! Personal use license - "
            "please do not resell or redistribute.").font = Font(italic=True, color=MUTED, size=9)

    # ---------------------------------------------------------------- Setup
    st = wb.create_sheet("Setup")
    base(st, {"A": 3, "B": 28, "C": 18, "D": 4, "E": 28, "F": 18})
    title(st, "Setup", "Edit the white cells. Everything else updates automatically.")
    cell(st, "B4", "Starting balance (Jan 1)", bold=True, fill=SOFT_FILL)
    cell(st, "C4", 2500, MONEY)
    cell(st, "B5", "Monthly savings target", bold=True, fill=SOFT_FILL)
    cell(st, "C5", 500, MONEY)

    header(st, 7, 2, ["Expense category", "Monthly budget"])
    EXP_FIRST = 8
    for i, (name, amt) in enumerate(EXPENSE_CATS):
        cell(st, f"B{EXP_FIRST + i}", name)
        cell(st, f"C{EXP_FIRST + i}", amt, MONEY)
    EXP_LAST = EXP_FIRST + len(EXPENSE_CATS) - 1
    cell(st, f"B{EXP_LAST + 1}", "TOTAL", bold=True, fill=SOFT_FILL)
    cell(st, f"C{EXP_LAST + 1}", f"=SUM(C{EXP_FIRST}:C{EXP_LAST})", MONEY, bold=True, fill=SOFT_FILL)

    header(st, 7, 5, ["Income category", "Expected / month"])
    INC_FIRST = 8
    for i, name in enumerate(INCOME_CATS):
        cell(st, f"E{INC_FIRST + i}", name)
        cell(st, f"F{INC_FIRST + i}", 4000 if i == 0 else 0, MONEY)
    INC_LAST = INC_FIRST + len(INCOME_CATS) - 1
    cell(st, f"E{INC_LAST + 1}", "TOTAL", bold=True, fill=SOFT_FILL)
    cell(st, f"F{INC_LAST + 1}", f"=SUM(F{INC_FIRST}:F{INC_LAST})", MONEY, bold=True, fill=SOFT_FILL)

    # Combined category list for the Transactions dropdown (income first, then expenses).
    ALL_ROWS = len(INCOME_CATS) + len(EXPENSE_CATS)
    st["H7"] = "All categories (auto)"
    st["H7"].font = Font(color=MUTED, size=9)
    for i in range(len(INCOME_CATS)):
        st[f"H{8 + i}"] = f"=E{INC_FIRST + i}"
    for i in range(len(EXPENSE_CATS)):
        st[f"H{8 + len(INCOME_CATS) + i}"] = f"=B{EXP_FIRST + i}"
    for i in range(ALL_ROWS):
        st[f"H{8 + i}"].font = Font(color=MUTED, size=9)
    st.column_dimensions["H"].width = 22
    CAT_LIST = f"Setup!$H$8:$H${7 + ALL_ROWS}"
    INC_RANGE = f"Setup!$E${INC_FIRST}:$E${INC_LAST}"

    # ---------------------------------------------------------------- Transactions
    tx = wb.create_sheet("Transactions")
    tx.sheet_properties.tabColor = ACCENT
    tx.sheet_view.showGridLines = False
    for col, w in {"A": 3, "B": 13, "C": 34, "D": 22, "E": 12, "F": 14, "G": 8}.items():
        tx.column_dimensions[col].width = w
    tx["B1"] = "Transactions"
    tx["B1"].font = Font(name="Georgia", size=22, bold=True, color=PRIMARY)
    tx["B2"] = "One row per income or expense. Enter amounts as positive numbers."
    tx["B2"].font = Font(italic=True, color=MUTED)
    header(tx, 4, 2, ["Date", "Description", "Category", "Type", "Amount", "Month"])
    TX_FIRST, TX_LAST = 5, 4 + TX_ROWS
    samples = [
        (f"{YEAR}-01-01", "January paycheck", "Salary", 4000),
        (f"{YEAR}-01-02", "Rent", "Rent / Mortgage", 1200),
        (f"{YEAR}-01-05", "Weekly groceries", "Groceries", 112.40),
        (f"{YEAR}-01-08", "Electric bill", "Utilities", 86.15),
        (f"{YEAR}-01-12", "Logo design gig", "Freelance", 350),
        (f"{YEAR}-01-15", "Dinner with friends", "Dining Out", 64.80),
    ]
    from datetime import date
    if demo:
        import random
        random.seed(7)
        samples = []
        for m in range(1, 13):
            samples.append((f"{YEAR}-{m:02d}-01", "Paycheck", "Salary", 4000))
            if m % 2 == 0:
                samples.append((f"{YEAR}-{m:02d}-10", "Freelance project", "Freelance", random.randint(200, 900)))
            for name, budget in EXPENSE_CATS:
                samples.append((f"{YEAR}-{m:02d}-{random.randint(2, 28):02d}", name, name,
                                round(budget * random.uniform(0.7, 1.25), 2)))
    for r in range(TX_FIRST, TX_LAST + 1):
        for col in "BCDEFG":
            c = tx[f"{col}{r}"]
            c.border = BORDER
            c.font = Font(color=INK)
        tx[f"B{r}"].number_format = "yyyy-mm-dd"
        tx[f"F{r}"].number_format = MONEY
        tx[f"E{r}"] = f'=IF(D{r}="","",IF(COUNTIF({INC_RANGE},D{r})>0,"Income","Expense"))'
        tx[f"G{r}"] = f'=IF(B{r}="","",MONTH(B{r}))'
        for col in "EG":
            tx[f"{col}{r}"].fill = SOFT_FILL
            tx[f"{col}{r}"].font = Font(color=MUTED)
    for i, (d, desc, cat, amt) in enumerate(samples):
        r = TX_FIRST + i
        tx[f"B{r}"] = date.fromisoformat(d)
        tx[f"C{r}"], tx[f"D{r}"], tx[f"F{r}"] = desc, cat, amt
    tx.freeze_panes = "B5"
    dv = DataValidation(type="list", formula1=f"={CAT_LIST}", allow_blank=True,
                        errorTitle="Unknown category", error="Pick a category from the list (edit them on Setup).")
    tx.add_data_validation(dv)
    dv.add(f"D{TX_FIRST}:D{TX_LAST}")
    dv_date = DataValidation(type="date", operator="greaterThan", formula1="1", allow_blank=True)
    tx.add_data_validation(dv_date)
    dv_date.add(f"B{TX_FIRST}:B{TX_LAST}")
    tx.conditional_formatting.add(
        f"B{TX_FIRST}:G{TX_LAST}",
        FormulaRule(formula=[f'$E{TX_FIRST}="Income"'], fill=PatternFill("solid", fgColor=PRIMARY_LIGHT)))

    TX_AMT = f"Transactions!$F${TX_FIRST}:$F${TX_LAST}"
    TX_CAT = f"Transactions!$D${TX_FIRST}:$D${TX_LAST}"
    TX_TYPE = f"Transactions!$E${TX_FIRST}:$E${TX_LAST}"
    TX_MON = f"Transactions!$G${TX_FIRST}:$G${TX_LAST}"

    # ---------------------------------------------------------------- Dashboard
    db = wb.create_sheet("Dashboard", 1)
    base(db, {"A": 3, "B": 22, "C": 17, "D": 17, "E": 17, "F": 17, "G": 17, "H": 4,
              "I": 22, "J": 14, "K": 14, "L": 14, "M": 12})
    title(db, f"{YEAR} Dashboard", "Updates automatically from your Transactions.")

    # KPI cards
    kpis = [
        ("B", "Total income", f'=SUMIFS({TX_AMT},{TX_TYPE},"Income")', MONEY),
        ("C", "Total spent", f'=SUMIFS({TX_AMT},{TX_TYPE},"Expense")', MONEY),
        ("D", "Net saved", "=B5-C5", MONEY),
        ("E", "Savings rate", '=IF(B5=0,0,D5/B5)', PCT),
        ("F", "Balance now", "=Setup!C4+D5", MONEY),
    ]
    for col, label, formula, fmt in kpis:
        cell(db, f"{col}4", label, bold=True, fill=HEADER_FILL, color=WHITE, align="center")
        c = cell(db, f"{col}5", formula, fmt, bold=True, fill=SOFT_FILL, align="center")
        c.font = Font(name="Georgia", size=13, bold=True, color=PRIMARY)
    db.row_dimensions[5].height = 30

    # Monthly table
    header(db, 7, 2, ["Month", "Income", "Expenses", "Net", "Savings rate", "Running balance"])
    for i, m in enumerate(MONTHS):
        r = 8 + i
        cell(db, f"B{r}", m, bold=True, fill=SOFT_FILL)
        cell(db, f"C{r}", f'=SUMIFS({TX_AMT},{TX_TYPE},"Income",{TX_MON},{i + 1})', MONEY, fill=SOFT_FILL)
        cell(db, f"D{r}", f'=SUMIFS({TX_AMT},{TX_TYPE},"Expense",{TX_MON},{i + 1})', MONEY, fill=SOFT_FILL)
        cell(db, f"E{r}", f"=C{r}-D{r}", MONEY, fill=SOFT_FILL)
        cell(db, f"F{r}", f"=IF(C{r}=0,0,E{r}/C{r})", PCT, fill=SOFT_FILL)
        prev = "Setup!$C$4" if i == 0 else f"G{r - 1}"
        cell(db, f"G{r}", f"={prev}+E{r}", MONEY, fill=SOFT_FILL)
    cell(db, "B20", "TOTAL", bold=True, fill=HEADER_FILL, color=WHITE)
    for col in "CDE":
        cell(db, f"{col}20", f"=SUM({col}8:{col}19)", MONEY, bold=True, fill=SOFT_FILL)
    cell(db, "F20", "=IF(C20=0,0,E20/C20)", PCT, bold=True, fill=SOFT_FILL)
    cell(db, "G20", "=G19", MONEY, bold=True, fill=SOFT_FILL)
    db.conditional_formatting.add("E8:E19", CellIsRule(operator="lessThan", formula=["0"],
                                  fill=PatternFill("solid", fgColor=RED_LIGHT), font=Font(color=RED)))

    # Category budget vs actual (year to date)
    header(db, 7, 9, ["Category", "Budget YTD", "Actual", "Remaining", "% used"])
    db["I6"] = "Budget YTD = monthly budget x months with activity"
    db["I6"].font = Font(size=9, italic=True, color=MUTED)
    months_active = f"MAX(1,MAX({TX_MON}))"
    for i in range(len(EXPENSE_CATS)):
        r = 8 + i
        s = EXP_FIRST + i
        cell(db, f"I{r}", f"=Setup!B{s}", fill=SOFT_FILL)
        cell(db, f"J{r}", f"=Setup!C{s}*{months_active}", MONEY, fill=SOFT_FILL)
        cell(db, f"K{r}", f"=SUMIFS({TX_AMT},{TX_CAT},I{r})", MONEY, fill=SOFT_FILL)
        cell(db, f"L{r}", f"=J{r}-K{r}", MONEY, fill=SOFT_FILL)
        cell(db, f"M{r}", f"=IF(J{r}=0,0,K{r}/J{r})", PCT, fill=SOFT_FILL)
    cat_last = 7 + len(EXPENSE_CATS)
    db.conditional_formatting.add(f"I8:M{cat_last}", FormulaRule(
        formula=["$M8>1"], fill=PatternFill("solid", fgColor=RED_LIGHT), font=Font(color=RED, bold=True)))
    db.conditional_formatting.add(f"M8:M{cat_last}", DataBarRule(
        start_type="num", start_value=0, end_type="num", end_value=1, color=BAR))

    bar = BarChart()
    bar.title = "Income vs. Expenses"
    bar.y_axis.title = None
    bar.height, bar.width = 7.5, 16
    bar.add_data(Reference(db, min_col=3, max_col=4, min_row=7, max_row=19), titles_from_data=True)
    bar.set_categories(Reference(db, min_col=2, min_row=8, max_row=19))
    bar.series[0].graphicalProperties.solidFill = PRIMARY
    bar.series[1].graphicalProperties.solidFill = ACCENT
    db.add_chart(bar, "B23")

    pie = PieChart()
    pie.title = "Where the money goes"
    pie.height, pie.width = 7.5, 12
    pie.add_data(Reference(db, min_col=11, min_row=7, max_row=cat_last), titles_from_data=True)
    pie.set_categories(Reference(db, min_col=9, min_row=8, max_row=cat_last))
    db.add_chart(pie, "I28")

    # ---------------------------------------------------------------- Savings Goals
    sg = wb.create_sheet("Savings Goals")
    base(sg, {"A": 3, "B": 26, "C": 14, "D": 14, "E": 13, "F": 14, "G": 14, "H": 16})
    title(sg, "Savings Goals", "Name a goal, set a target and a deadline - see what to save each month.")
    header(sg, 4, 2, ["Goal", "Target", "Saved so far", "Progress", "Remaining", "Deadline", "Save per month"])
    goals = [("Emergency fund", 6000, 1500, f"{YEAR}-12-31"), ("Vacation", 2500, 400, f"{YEAR}-07-01"),
             ("New laptop", 1400, 200, f"{YEAR}-05-01")]
    for i in range(12):
        r = 5 + i
        for col in "BCDG":
            cell(sg, f"{col}{r}")
        sg[f"C{r}"].number_format = sg[f"D{r}"].number_format = MONEY
        sg[f"G{r}"].number_format = "yyyy-mm-dd"
        if i < len(goals):
            g = goals[i]
            sg[f"B{r}"], sg[f"C{r}"], sg[f"D{r}"] = g[0], g[1], g[2]
            sg[f"G{r}"] = date.fromisoformat(g[3])
        cell(sg, f"E{r}", f'=IF(C{r}="","",MIN(1,D{r}/C{r}))', PCT, fill=SOFT_FILL)
        cell(sg, f"F{r}", f'=IF(C{r}="","",MAX(0,C{r}-D{r}))', MONEY, fill=SOFT_FILL)
        cell(sg, f"H{r}", f'=IF(OR(C{r}="",G{r}=""),"",IF(F{r}=0,0,'
                          f'F{r}/MAX(1,DATEDIF(TODAY(),G{r},"m")+(G{r}<=TODAY()))))', MONEY, fill=SOFT_FILL)
    sg.conditional_formatting.add("E5:E16", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                       end_value=1, color=BAR))

    # ---------------------------------------------------------------- Debt Payoff
    dp = wb.create_sheet("Debt Payoff")
    base(dp, {"A": 3, "B": 24, "C": 14, "D": 11, "E": 14, "F": 14, "G": 15, "H": 16, "I": 16})
    title(dp, "Debt Payoff Planner", "Add extra payments and watch your debt-free date move closer.")
    header(dp, 4, 2, ["Debt", "Balance", "APR", "Min payment", "Extra / month",
                      "Months to payoff", "Debt-free date", "Total interest"])
    debts = [("Credit card", 3200, 0.229, 95, 100), ("Car loan", 9800, 0.069, 280, 0),
             ("Student loan", 14500, 0.045, 160, 0)]
    for i in range(10):
        r = 5 + i
        for col in "BCDEF":
            cell(dp, f"{col}{r}")
        dp[f"C{r}"].number_format = dp[f"E{r}"].number_format = dp[f"F{r}"].number_format = MONEY
        dp[f"D{r}"].number_format = "0.0%"
        if i < len(debts):
            for col, v in zip("BCDEF", debts[i]):
                dp[f"{col}{r}"] = v
        pay = f"(E{r}+F{r})"
        cell(dp, f"G{r}", f'=IF(OR(C{r}="",C{r}=0),"",IF(D{r}=0,ROUNDUP(C{r}/{pay},0),'
                          f'IF({pay}<=C{r}*D{r}/12,"Payment too low",ROUNDUP(NPER(D{r}/12,-{pay},C{r}),0))))',
             "0", fill=SOFT_FILL, align="center")
        cell(dp, f"H{r}", f'=IF(ISNUMBER(G{r}),EDATE(TODAY(),G{r}),"")', "mmm yyyy", fill=SOFT_FILL, align="center")
        cell(dp, f"I{r}", f'=IF(ISNUMBER(G{r}),IF(D{r}=0,0,MAX(0,{pay}*NPER(D{r}/12,-{pay},C{r})-C{r})),"")',
             MONEY, fill=SOFT_FILL)
    cell(dp, "B15", "TOTAL", bold=True, fill=HEADER_FILL, color=WHITE)
    for col in "CEF":
        cell(dp, f"{col}15", f"=SUM({col}5:{col}14)", MONEY, bold=True, fill=SOFT_FILL)
    cell(dp, "G15", "=IF(COUNT(G5:G14)=0,\"\",MAX(G5:G14))", "0", bold=True, fill=SOFT_FILL, align="center")
    cell(dp, "H15", '=IF(ISNUMBER(G15),EDATE(TODAY(),G15),"")', "mmm yyyy", bold=True, fill=SOFT_FILL, align="center")
    cell(dp, "I15", "=SUM(I5:I14)", MONEY, bold=True, fill=SOFT_FILL)
    dp["B17"] = "Tip: put any extra money on the highest-APR debt first (avalanche) to pay the least interest."
    dp["B17"].font = Font(italic=True, color=MUTED)

    # ---------------------------------------------------------------- Bill Calendar
    bc = wb.create_sheet("Bill Tracker")
    base(bc, {"A": 3, "B": 24, "C": 12, "D": 13, **{chr(ord("E") + i): 6 for i in range(12)}})
    title(bc, "Bill Tracker", "List recurring bills, then type x in each month once paid.")
    header(bc, 4, 2, ["Bill", "Due day", "Amount"] + MONTHS)
    bills = [("Rent", 1, 1200), ("Electric", 15, 90), ("Phone", 20, 55), ("Streaming", 3, 15.99)]
    for i in range(15):
        r = 5 + i
        for col_idx in range(2, 17):
            c = bc.cell(row=r, column=col_idx)
            c.border, c.fill = BORDER, INPUT_FILL
            if col_idx >= 5:
                c.alignment = Alignment(horizontal="center")
        bc[f"D{r}"].number_format = MONEY
        if i < len(bills):
            bc[f"B{r}"], bc[f"C{r}"], bc[f"D{r}"] = bills[i]
    bc.conditional_formatting.add("E5:P19", CellIsRule(operator="equal", formula=['"x"'],
                                  fill=PatternFill("solid", fgColor=BAR), font=Font(bold=True, color=WHITE)))

    print_areas = {"Start Here": "A1:B22", "Setup": "A1:H32", "Dashboard": "A1:M43",
                   "Transactions": "A1:G60", "Savings Goals": "A1:H17", "Debt Payoff": "A1:I17",
                   "Bill Tracker": "A1:P20"}
    for sheet in wb:
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1 if sheet.title != "Transactions" else 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_area = print_areas[sheet.title]

    wb.active = 0
    wb.save(out)
    return out


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if a != "--demo"]
    print("Saved", build(args[0] if args else "product/2027-Budget-Finance-Planner.xlsx", demo="--demo" in sys.argv))
